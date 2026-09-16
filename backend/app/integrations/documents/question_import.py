"""Deterministic, rule-based extraction of assessment questions from an
uploaded PDF, DOCX, XLSX, or CSV file (docs/assessments.md's "Import
Questions" workflow). No AI/LLM involved — this never guesses at a correct
answer it can't find explicitly in the document (CLAUDE.md § 5: no fake
implementations); a question whose correct answer can't be determined is
skipped and surfaced as a warning rather than silently invented.

Callers get back plain data (not persisted) so the recruiter can preview,
edit, reorder, and select questions before anything is saved — the actual
persistence still goes through the existing `assessment_service.create_assessment`
/ `AssessmentCreateRequest` path, so this module never touches the database.

Supported conventions:
  - CSV / XLSX: a header row with a `question`/`prompt` column, one or more
    `option`/`option_N` (or `a`, `b`, `c`, `d`, ...) columns, and a
    `correct`/`answer` column identifying the right option(s) by letter,
    1-based index, or exact text match. Optional `type` and `points`
    columns.
  - PDF / DOCX: numbered questions ("1. ...") followed by lettered options
    ("A) ...", "B) ..."), with the correct option(s) marked either by a
    trailing "*" on the option line or an explicit "Answer: B" /
    "Correct: B, C" line.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.exceptions import AppError


@dataclass
class ParsedOption:
    label: str
    is_correct: bool = False


@dataclass
class ParsedQuestion:
    prompt: str
    type: str  # "MCQ_SINGLE" | "MCQ_MULTI"
    points: int = 1
    options: list[ParsedOption] = field(default_factory=list)


@dataclass
class QuestionImportResult:
    questions: list[ParsedQuestion]
    warnings: list[str]


class QuestionImportError(AppError):
    code = "question_import_failed"


def parse_questions(*, content: bytes, filename: str) -> QuestionImportResult:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if suffix == "csv":
        rows = _read_csv_rows(content)
        return _parse_tabular(rows)
    if suffix == "xlsx":
        rows = _read_xlsx_rows(content)
        return _parse_tabular(rows)
    if suffix == "docx":
        text = _read_docx_text(content)
        return _parse_freeform_text(text)
    if suffix == "pdf":
        text = _read_pdf_text(content)
        return _parse_freeform_text(text)

    raise QuestionImportError(
        f"Unsupported file type: .{suffix or 'unknown'}. Upload a PDF, DOCX, XLSX, or CSV file."
    )


# --- File readers ---------------------------------------------------------


def _read_csv_rows(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise QuestionImportError(f"Could not read this CSV file as text: {exc}") from exc
    reader = csv.reader(io.StringIO(text))
    return [row for row in reader if any(cell.strip() for cell in row)]


def _read_xlsx_rows(content: bytes) -> list[list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        rows = [
            ["" if cell is None else str(cell) for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    except Exception as exc:
        raise QuestionImportError(f"Could not read this XLSX file: {exc}") from exc
    return [row for row in rows if any(cell.strip() for cell in row)]


def _read_docx_text(content: bytes) -> str:
    try:
        document = Document(io.BytesIO(content))
        lines = [paragraph.text for paragraph in document.paragraphs]
    except Exception as exc:
        raise QuestionImportError(f"Could not read this DOCX file: {exc}") from exc
    text = "\n".join(lines).strip()
    if not text:
        raise QuestionImportError("No extractable text found in this DOCX file.")
    return text


def _read_pdf_text(content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise QuestionImportError(f"Could not read this PDF file: {exc}") from exc
    text = text.strip()
    if not text:
        raise QuestionImportError(
            "No extractable text found in this PDF (it may be a scanned image)."
        )
    return text


# --- Tabular (CSV/XLSX) parsing -------------------------------------------

_OPTION_COLUMN_RE = re.compile(r"^option[_ ]?(\d+)$")
_LETTER_OPTION_COLUMNS = ["a", "b", "c", "d", "e", "f"]


def _normalize_header(raw: str) -> str:
    return re.sub(r"[\s\-]+", "_", raw.strip().lower())


def _parse_tabular(rows: list[list[str]]) -> QuestionImportResult:
    if not rows:
        raise QuestionImportError("This file has no rows.")

    header = [_normalize_header(cell) for cell in rows[0]]
    question_idx = next((i for i, h in enumerate(header) if h in ("question", "prompt")), None)
    if question_idx is None:
        raise QuestionImportError(
            "Expected a 'question' column in the first row. "
            "See the import template for the expected format."
        )

    option_cols: list[tuple[int, int]] = []  # (column index, option number)
    for i, h in enumerate(header):
        match = _OPTION_COLUMN_RE.match(h)
        if match:
            option_cols.append((i, int(match.group(1))))
    if not option_cols:
        for i, h in enumerate(header):
            if h in _LETTER_OPTION_COLUMNS:
                option_cols.append((i, _LETTER_OPTION_COLUMNS.index(h) + 1))
    option_cols.sort(key=lambda pair: pair[1])

    correct_idx = next(
        (i for i, h in enumerate(header) if h in ("correct", "answer", "correct_answer")), None
    )
    type_idx = next((i for i, h in enumerate(header) if h == "type"), None)
    points_idx = next((i for i, h in enumerate(header) if h == "points"), None)

    if not option_cols:
        raise QuestionImportError(
            "Expected option columns such as 'option_1', 'option_2', ... (or 'a', 'b', 'c', ...)."
        )

    questions: list[ParsedQuestion] = []
    warnings: list[str] = []

    for row_number, row in enumerate(rows[1:], start=2):
        cell = lambda i: row[i].strip() if i < len(row) and row[i] is not None else ""  # noqa: E731

        prompt = cell(question_idx)
        if not prompt:
            continue

        options: list[str] = []
        for col_idx, _ in option_cols:
            value = cell(col_idx)
            if value:
                options.append(value)
        if len(options) < 2:
            warnings.append(f"Row {row_number}: skipped — fewer than 2 non-empty options.")
            continue

        correct_indices = _resolve_correct_indices(
            raw=cell(correct_idx) if correct_idx is not None else "",
            options=options,
        )
        if not correct_indices:
            warnings.append(
                f"Row {row_number}: skipped — could not determine the correct answer."
            )
            continue

        declared_type = cell(type_idx).lower() if type_idx is not None else ""
        if "multi" in declared_type:
            q_type = "MCQ_MULTI"
        elif "single" in declared_type:
            q_type = "MCQ_SINGLE"
        else:
            q_type = "MCQ_MULTI" if len(correct_indices) > 1 else "MCQ_SINGLE"

        points_raw = cell(points_idx) if points_idx is not None else ""
        points = int(points_raw) if points_raw.isdigit() else 1

        questions.append(
            ParsedQuestion(
                prompt=prompt,
                type=q_type,
                points=points,
                options=[
                    ParsedOption(label=label, is_correct=(i in correct_indices))
                    for i, label in enumerate(options)
                ],
            )
        )

    return QuestionImportResult(questions=questions, warnings=warnings)


def _resolve_correct_indices(*, raw: str, options: list[str]) -> set[int]:
    """`raw` may name correct options by 1-based index ("2"), by letter
    ("B"), or by exact text match — comma/semicolon/'and'-separated for
    multiple answers."""
    if not raw:
        return set()

    tokens = [t.strip() for t in re.split(r"[,;/]| and ", raw, flags=re.IGNORECASE) if t.strip()]
    indices: set[int] = set()
    for token in tokens:
        if token.isdigit():
            i = int(token) - 1
            if 0 <= i < len(options):
                indices.add(i)
            continue
        if len(token) == 1 and token.isalpha():
            i = ord(token.upper()) - ord("A")
            if 0 <= i < len(options):
                indices.add(i)
                continue
        for i, label in enumerate(options):
            if label.strip().lower() == token.lower():
                indices.add(i)
                break
    return indices


# --- Free-text (PDF/DOCX) parsing -----------------------------------------

_QUESTION_START_RE = re.compile(r"^\s*(\d+)[.)]\s+(.*\S)\s*$")
_OPTION_LINE_RE = re.compile(r"^\s*\(?([A-Za-z])[.)]\s+(.*\S)\s*$")
_ANSWER_LINE_RE = re.compile(r"^\s*(?:answer|correct)s?\s*[:\-]\s*(.+\S)\s*$", re.IGNORECASE)
_POINTS_SUFFIX_RE = re.compile(r"[\[(]\s*(\d+)\s*(?:pts?|points?)\s*[\])]", re.IGNORECASE)


def _parse_freeform_text(text: str) -> QuestionImportResult:
    lines = [line.rstrip() for line in text.splitlines()]

    blocks: list[list[str]] = []
    current: list[str] | None = None
    for line in lines:
        if _QUESTION_START_RE.match(line):
            if current is not None:
                blocks.append(current)
            current = [line]
        elif current is not None:
            current.append(line)
    if current is not None:
        blocks.append(current)

    if not blocks:
        raise QuestionImportError(
            "No numbered questions (e.g. '1. What is ...') were found in this document. "
            "See the import template for the expected format."
        )

    questions: list[ParsedQuestion] = []
    warnings: list[str] = []

    for block_number, block in enumerate(blocks, start=1):
        start_match = _QUESTION_START_RE.match(block[0])
        assert start_match is not None
        prompt_lines = [start_match.group(2)]

        options: list[ParsedOption] = []
        answer_letters: set[str] | None = None
        in_options = False

        for line in block[1:]:
            if not line.strip():
                continue
            option_match = _OPTION_LINE_RE.match(line)
            answer_match = _ANSWER_LINE_RE.match(line)
            if answer_match:
                answer_letters = {
                    t.strip().upper()
                    for t in re.split(r"[,;/]| and ", answer_match.group(1), flags=re.IGNORECASE)
                    if t.strip()
                }
                continue
            if option_match:
                in_options = True
                label = option_match.group(2).strip()
                is_marked = label.endswith("*")
                if is_marked:
                    label = label[:-1].strip()
                options.append(ParsedOption(label=label, is_correct=is_marked))
                continue
            if not in_options:
                prompt_lines.append(line.strip())

        prompt = " ".join(p for p in prompt_lines if p).strip()
        points_match = _POINTS_SUFFIX_RE.search(prompt)
        points = int(points_match.group(1)) if points_match else 1
        if points_match:
            prompt = _POINTS_SUFFIX_RE.sub("", prompt).strip()

        if len(options) < 2:
            warnings.append(f"Question {block_number}: skipped — fewer than 2 options found.")
            continue

        if answer_letters:
            for option in options:
                option.is_correct = False
            for i, option in enumerate(options):
                letter = chr(ord("A") + i)
                if letter in answer_letters:
                    option.is_correct = True

        correct_count = sum(1 for o in options if o.is_correct)
        if correct_count == 0:
            warnings.append(
                f"Question {block_number}: skipped — no correct answer marked "
                "(mark it with a trailing '*' or an 'Answer: X' line)."
            )
            continue

        questions.append(
            ParsedQuestion(
                prompt=prompt,
                type="MCQ_MULTI" if correct_count > 1 else "MCQ_SINGLE",
                points=points,
                options=options,
            )
        )

    return QuestionImportResult(questions=questions, warnings=warnings)
