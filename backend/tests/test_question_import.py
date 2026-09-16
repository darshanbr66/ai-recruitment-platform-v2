"""Unit tests for the rule-based question-import parser
(app/integrations/documents/question_import.py) — pure functions, no DB.
Covers the "Import Questions" workflow's core promise: never fabricate a
correct answer it can't find explicitly in the document."""

import io

import pytest
from docx import Document
from openpyxl import Workbook

from app.integrations.documents.question_import import QuestionImportError, parse_questions


def test_csv_single_and_multi_answer() -> None:
    csv_bytes = (
        b"question,option_1,option_2,option_3,option_4,correct,type\n"
        b"What is 2+2?,3,4,5,6,2,single\n"
        b'Pick prime numbers,2,4,6,7,"1,4",multiple\n'
    )
    result = parse_questions(content=csv_bytes, filename="questions.csv")

    assert result.warnings == []
    assert len(result.questions) == 2

    q1 = result.questions[0]
    assert q1.prompt == "What is 2+2?"
    assert q1.type == "MCQ_SINGLE"
    assert [o.is_correct for o in q1.options] == [False, True, False, False]

    q2 = result.questions[1]
    assert q2.type == "MCQ_MULTI"
    assert [o.is_correct for o in q2.options] == [True, False, False, True]


def test_csv_letter_answer_key() -> None:
    csv_bytes = b"question,option_1,option_2,option_3,correct\nCapital of France?,Berlin,Paris,Rome,B\n"
    result = parse_questions(content=csv_bytes, filename="q.csv")
    assert result.questions[0].options[1].label == "Paris"
    assert result.questions[0].options[1].is_correct is True


def test_csv_row_without_resolvable_answer_is_skipped_not_guessed() -> None:
    csv_bytes = b"question,option_1,option_2,correct\nUnanswerable?,A,B,\n"
    result = parse_questions(content=csv_bytes, filename="q.csv")
    assert result.questions == []
    assert "skipped" in result.warnings[0].lower()


def test_csv_missing_question_column_is_a_clear_error() -> None:
    csv_bytes = b"a,b,c\n1,2,3\n"
    with pytest.raises(QuestionImportError, match="question"):
        parse_questions(content=csv_bytes, filename="q.csv")


def test_xlsx_roundtrip() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["question", "option_1", "option_2", "option_3", "correct"])
    sheet.append(["2 + 2 = ?", "3", "4", "5", "2"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = parse_questions(content=buffer.getvalue(), filename="questions.xlsx")
    assert len(result.questions) == 1
    assert result.questions[0].options[1].is_correct is True


def test_docx_numbered_questions_with_asterisk_and_answer_line() -> None:
    document = Document()
    document.add_paragraph("1. What is the capital of France?")
    document.add_paragraph("A) Berlin")
    document.add_paragraph("B) Paris*")
    document.add_paragraph("C) Madrid")
    document.add_paragraph("2. Which are prime numbers? [2 pts]")
    document.add_paragraph("A) 2")
    document.add_paragraph("B) 4")
    document.add_paragraph("C) 5")
    document.add_paragraph("D) 9")
    document.add_paragraph("Answer: A, C")
    document.add_paragraph("3. No answer marked at all")
    document.add_paragraph("A) x")
    document.add_paragraph("B) y")
    buffer = io.BytesIO()
    document.save(buffer)

    result = parse_questions(content=buffer.getvalue(), filename="questions.docx")

    assert len(result.questions) == 2
    q1 = result.questions[0]
    assert q1.prompt == "What is the capital of France?"
    assert q1.type == "MCQ_SINGLE"
    assert q1.options[1].label == "Paris"
    assert q1.options[1].is_correct is True

    q2 = result.questions[1]
    assert q2.type == "MCQ_MULTI"
    assert q2.points == 2
    assert "[2 pts]" not in q2.prompt
    assert [o.is_correct for o in q2.options] == [True, False, True, False]

    assert len(result.warnings) == 1
    assert "Question 3" in result.warnings[0]


def test_docx_with_no_numbered_questions_is_a_clear_error() -> None:
    document = Document()
    document.add_paragraph("This document has no numbered questions in it.")
    buffer = io.BytesIO()
    document.save(buffer)

    with pytest.raises(QuestionImportError, match="numbered questions"):
        parse_questions(content=buffer.getvalue(), filename="q.docx")


def test_unsupported_file_type_is_rejected() -> None:
    with pytest.raises(QuestionImportError, match="Unsupported file type"):
        parse_questions(content=b"whatever", filename="questions.txt")
